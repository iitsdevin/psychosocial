#!/usr/bin/env python3
"""Build the Principal Wellbeing funding email response builder workbook.

Generates an interactive Excel tool that triages principal wellbeing funding
queries using the six decision-making framework questions (plus panel and
travel routing questions) and assembles a draft email reply in the team's
house style. Styled to match the Statewide Services interactive list of
support (WA Department of Education magenta/teal, hidden gridlines,
instruction panel, coloured tabs).

Usage: python3 build_pwi_email_tool.py [output.xlsx]
"""

import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- palette --
MAGENTA = "9E0E6C"          # WA Education magenta (Statewide primary)
TEAL = "00809E"             # Statewide secondary
MAGENTA_TINT = "F9E6F1"
TEAL_TINT = "E0F2F6"
INPUT_YELLOW = "FFF2CC"
GREY_TEXT = "595959"
DRAFT_BG = "F7F7F7"

GREEN_FILL, GREEN_FONT = "C6EFCE", "006100"
RED_FILL, RED_FONT = "FFC7CE", "9C0006"
AMBER_FILL, AMBER_FONT = "FFEB9C", "9C6500"

FONT = "Arial"
NL2 = 'CHAR(10)&CHAR(10)'

thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(color):
    return PatternFill("solid", fgColor=color)


def banner(ws, cell_range, text, bg, size=16, color="FFFFFF", bold=True):
    ws.merge_cells(cell_range)
    c = ws[cell_range.split(":")[0]]
    c.value = text
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    c.fill = fill(bg)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    first, last = cell_range.split(":")
    row = int("".join(ch for ch in first if ch.isdigit()))
    for col in range(ws[first].column, ws[last].column + 1):
        ws.cell(row=row, column=col).fill = fill(bg)


def step_header(ws, cell_range, text):
    banner(ws, cell_range, text, MAGENTA, size=11)


def note(ws, coord, text, size=9):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, size=size, italic=True, color=GREY_TEXT)
    c.alignment = Alignment(vertical="center", wrap_text=True)


def label(ws, coord, text, bold=True):
    c = ws[coord]
    c.value = text
    c.font = Font(name=FONT, size=10, bold=bold)
    c.alignment = Alignment(vertical="center", wrap_text=True)


def input_cell(ws, coord, value=""):
    c = ws[coord]
    c.value = value
    c.font = Font(name=FONT, size=10)
    c.fill = fill(INPUT_YELLOW)
    c.border = box
    c.alignment = Alignment(vertical="center")
    return c


def no_gridlines(ws):
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------ content --
QUESTIONS = [
    ("Q1", "Will this directly benefit the principal's health, wellbeing or both?"),
    ("Q2", "Does this align with the priorities in the health and wellbeing strategies or panel service categories?"),
    ("Q3", "Would this expenditure uphold Department values and standards, as per the Integrity framework and Code of Conduct?"),
    ("Q4", "Is the provider reputable, appropriately insured and qualified to deliver the service?"),
    ("Q5", "Does this service meet a need that cannot be provided for by current Department programs or services?"),
    ("Q6", "Have all supplier-related conflicts of interest been identified, reported and appropriately managed (or do none apply)?"),
]
ROUTING = [
    ("Panel", "Does this align with the panel of wellbeing service providers?"),
    ("Travel", "Does this involve travel or accommodation?"),
]

LIBRARY = [
    ("Q1", "Yes", "The proposed expenditure meets the first test of the decision-making framework: it is intended to directly benefit your health, wellbeing or both as a principal."),
    ("Q1", "No", "The proposed expenditure does not appear to directly benefit your health or wellbeing as a principal. The targeted funding is to address your needs as a principal - it is not intended for staff wellbeing or general operating costs - so we recommend reconsidering this expenditure."),
    ("Q1", "Maybe", "It is not clear that this would directly benefit your health or wellbeing as a principal. We would be happy to talk this through with you before you proceed."),
    ("Q2", "Yes", "It aligns with the priorities in the health and wellbeing strategies and the panel service categories (executive and leadership coaching, executive health assessment, and mental health support)."),
    ("Q2", "No", "It does not align with the priorities in the health and wellbeing strategies or the panel service categories, so we recommend reconsidering this expenditure."),
    ("Q2", "Maybe", "The alignment with the health and wellbeing strategies or the panel service categories is not clear - please check with us before proceeding."),
    ("Q3", "Yes", "The expenditure would uphold Department values and standards, in line with our Integrity framework and Code of Conduct (including Standard 8 - use public resources responsibly)."),
    ("Q3", "No", "The expenditure does not appear to uphold Department values and standards under the Integrity framework and Code of Conduct, so we recommend reconsidering it."),
    ("Q3", "Maybe", "It is not clear that the expenditure would uphold Department values and standards under the Integrity framework and Code of Conduct - please contact us before proceeding."),
    ("Q4", "Yes", "The provider is reputable, appropriately insured and qualified to deliver the service."),
    ("Q4", "No", "The provider's qualifications, insurance or reputation cannot be confirmed, so please do not proceed until this is resolved. Information about managing provider risks is available in the Develop a staff health and wellbeing plan Ikon service."),
    ("Q4", "Maybe", "The provider's qualifications, insurance or reputation have not yet been confirmed. Where relevant, check the Australian Health Practitioner Regulation Agency (AHPRA) register and confirm the provider's insurance before proceeding."),
    ("Q5", "Yes", "The service meets a need that current Department programs or services do not already provide for."),
    ("Q5", "No", "Current Department programs or services already meet this need - for example, the Employee Assistance Program now provides up to 12 confidential counselling sessions per year. We recommend using the existing service."),
    ("Q5", "Maybe", "It is not clear whether current Department programs or services already meet this need. Please check the available services - including the Employee Assistance Program (up to 12 sessions per year) - before proceeding."),
    ("Q6", "Yes", "All supplier-related conflicts of interest have been identified, reported and appropriately managed (or none apply)."),
    ("Q6", "No", "A supplier-related conflict of interest has not yet been reported or managed. Please report the conflict of interest before proceeding."),
    ("Q6", "Maybe", "A possible supplier-related conflict of interest has not been confirmed as managed. Please report and resolve this before proceeding."),
    ("Panel", "Yes", "This service falls within the panel of wellbeing service providers' categories. Once the panel is live, we encourage engaging a provider through the panel; in the meantime, the principal wellbeing consultant can help you assess and access suitable services."),
    ("Panel", "No", "This service sits outside the panel of wellbeing service providers. You can engage the provider through local procurement, following the rules for buying goods and services. For purchases under $50,000, the Buy goods and services under $50,000 Ikon service applies, and purchases of less than $5,000 are low-value procurements provided value for money has been considered."),
    ("Panel", "Maybe", "It is not yet clear whether this service aligns with the panel of wellbeing service providers. The panel is non-mandatory and local procurement remains available - please check with us if you are unsure."),
    ("Travel", "Yes", "As the query involves travel or accommodation: the funding may be used for travel and accommodation where it is required to support access to an eligible principal wellbeing activity, in line with the Official Travel Policy and the associated Official Air Travel and Travel Accommodation, Meal and Other Expenses procedures. Travellers cannot approve their own travel - please seek approval from an Assistant Director of Education or above prior to booking. The funding is not intended to be used solely for travel or accommodation."),
    ("Travel", "No", ""),
    ("Travel", "Maybe", "If travel or accommodation ends up being involved, it may be funded where it supports the eligible wellbeing activity, in line with the Official Travel Policy - please check the approval requirements with us before booking."),
]

P = "\n\n"
PRESETS = [
    ("Treatments and therapies (e.g. physiotherapy) - not in scope",
     "To answer briefly: no - personal health care, treatments and therapies (including physiotherapy and chiropractic care) are not within the scope of the targeted wellbeing funding."
     + P + "To answer more thoroughly:"
     + P + "The Workers' Compensation and Injury Management Act 2023 (WA) prohibits employers from diverting workers away from making a claim, so wellbeing funds cannot be used for injury-related assessment or treatment services. Treatment for work-related pain or injury may be available through the workers' compensation process, and may include payment for expenses, loss of earnings or both. Personal medical expenses are considered a fringe benefit."
     + P + "While this service falls outside the scope of the funding, there are a range of other services you may wish to consider, including professional learning and resilience training, mental health support or pastoral care, coaching or mentoring (including executive and wellbeing coaching), and executive health and wellbeing assessments."),
    ("Memberships, gym, goods or devices - not in scope",
     "To answer briefly: no - memberships (such as gym or Pilates memberships), goods and personal devices are for the personal benefit of the employee, are considered a fringe benefit, and are not within the scope of the targeted wellbeing funding."
     + P + "You may instead wish to consider services such as professional learning and resilience training, mental health support or pastoral care, coaching or mentoring (including executive and wellbeing coaching), or an executive health and wellbeing assessment."),
    ("Travel and accommodation - in scope with conditions",
     "To answer briefly: yes - the wellbeing funding can be used toward travel and accommodation needed to support access to an eligible wellbeing activity (such as professional learning, a professional conference or similar)."
     + P + "To answer more thoroughly:"
     + P + "The funding may be used for travel and accommodation where the travel is required to support the wellbeing activity. It is not intended to be used solely for travel or accommodation."
     + P + "Travel needs to align with the Official Travel Policy and the associated Official Air Travel and Travel Accommodation, Meal and Other Expenses procedures. Travellers cannot approve their own travel: per the Travel Approval Schedule, please seek approval from an Assistant Director of Education or above prior to booking. For domestic air travel, the domestic air travel application form must be completed and approved before booking."),
    ("Conference or professional learning - in scope with conditions",
     "To answer briefly: yes - conferences and professional learning are considered appropriate expenditure where the majority of the program focuses on enhancing leadership capability and includes wellbeing topics for executives."
     + P + "To answer more thoroughly:"
     + P + "Suitable themes include work-related stress reduction, managing workload and role overload, resilience skills for leaders, decision-making under stress, and high-performance leadership without burnout."
     + P + "The targeted funding may complement your school's budget allocation for the principal's professional learning, and relevant travel costs, but it should not replace it. If travel is required to attend, it needs to align with the Official Travel Policy and associated procedures, with approval from an Assistant Director of Education or above prior to booking."),
    ("Pooling funding with other schools - supported",
     "To answer briefly: yes - you may choose to pool your wellbeing funding with other schools, such as those in your school network, to access higher cost support services (for example, leadership or wellbeing coaching)."
     + P + "To answer more thoroughly:"
     + P + "The decision-making framework applies to pooled funding in the same way as individual funding. You can choose an appropriate budget account or accounts to record the transfer of funds when pooling resources."
     + P + "For purchases under $50,000, follow the Buy goods and services under $50,000 Ikon service - the Verbal quotation template or Very simple purchase template may be useful. Purchases of less than $5,000 are considered low-value procurements and a panel is not required, provided value for money has been considered."),
    ("Recording purchases and acquittal - guidance",
     "To answer briefly: expenditure only needs to be lodged through Oracle, and you can choose an appropriate budget account or accounts to record purchases against."
     + P + "We do not monitor principal wellbeing expenditure, and you do not need to provide formal acquittal reports. We do encourage you to keep a record of your decisions to demonstrate accountable and ethical decision-making - the Manage school accounting records Ikon service provides guidance on record keeping."),
    ("Ergonomic assessment or equipment - school budget",
     "To answer briefly: ergonomic assessments and equipment are provided for in existing school funding, so they sit outside the targeted wellbeing funding."
     + P + "For guidance, refer to the Manage workstation ergonomics and Work comfortably at your computer Ikon services."),
    ("Catering - in scope with conditions",
     "To answer briefly: the funding may be used for catering where the expenditure complies with the Expenditure on Hospitality procedures and is related to a principal wellbeing activity."
     + P + "Please note that catering alone is not a principal wellbeing activity - it needs to be connected to an eligible wellbeing activity."),
]

LINKS = [
    ("Spend principal wellbeing targeted funding (decision-making framework)", "https://ikon.education.wa.edu.au/-/spend-principal-wellbeing-targeted-funding"),
    ("Examples of using the decision-making framework", "https://ikon.education.wa.edu.au/-/examples-of-using-the-decision-making-framework-for-wellbeing-expenditure-decisions/"),
    ("Integrity framework", "https://ikon.education.wa.edu.au/-/integrity-framework/"),
    ("Our Code of Conduct and standards", "https://ikon.education.wa.edu.au/-/our-code-of-conduct-and-standards/"),
    ("Report a conflict of interest", "https://ikon.education.wa.edu.au/-/declare-a-conflict-of-interest"),
    ("Access the Employee Assistance Program", "https://ikon.education.wa.edu.au/-/access-the-employee-assistance-program"),
    ("AHPRA practitioner register", "https://www.ahpra.gov.au/"),
    ("Buy goods and services under $50,000", "https://ikon.education.wa.edu.au/-/buy-goods-and-services-under-50-000"),
    ("Follow the rules for buying goods and services", "https://ikon.education.wa.edu.au/-/follow-rules-for-buying-goods-and-services"),
    ("Official Travel Policy", "https://www.education.wa.edu.au/web/policies/-/official-travel-policy"),
    ("Official Air Travel procedures", "https://www.education.wa.edu.au/web/policies/-/official-air-travel-procedures"),
    ("Travel Accommodation, Meal and Other Expenses procedures", "https://www.education.wa.edu.au/web/policies/-/travel-accommodation-meal-and-other-expenses-procedures"),
    ("Manage school accounting records", "https://ikon.education.wa.edu.au/-/manage-accounting-records"),
    ("Principal Wellbeing Connect Community", "https://connect.det.wa.edu.au/group/staff/ui/community/summary?coisp=DomainCommunity:6697628811"),
    ("Tenders WA (panel tender information for providers)", "https://www.tenders.wa.gov.au/"),
]

BUILD_OPTION = "Build from the framework questions"
CONTACT = ("If you have any further questions, please reach out again. You can contact "
           "Principal wellbeing expenditure support by phone on 9264 4201 or by email at "
           "whsw.staffwellbeing@education.wa.edu.au.")
PLACEHOLDER = ("[Pick a common response in Step 2, or answer the framework questions in "
               "Step 3, to build the body of your email.]")


def build(path):
    wb = Workbook()

    # ================================================== Start here ==
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = MAGENTA
    no_gridlines(ws)
    widths = {"A": 2, "B": 4, "C": 100, "D": 2}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.row_dimensions[2].height = 34
    banner(ws, "B2:C2", "Principal wellbeing funding - email response builder", MAGENTA, size=18)
    ws.row_dimensions[3].height = 20
    banner(ws, "B3:C3", "Triage a query with the decision-making framework and generate a draft reply", TEAL, size=11, bold=False)

    step_header(ws, "B5:C5", "How to use this tool")
    steps = [
        "1.  Open the 'Email builder' tab.",
        "2.  Step 1 - choose a greeting and add the recipient's name and the topic of their query.",
        "3.  Step 2 - if it is a recurring query (physiotherapy, travel, pooling funds, etc.), pick a common response from the dropdown. Done!",
        "4.  Step 3 - otherwise, answer the six decision-making framework questions (plus the two routing questions) using the Yes / No / Maybe dropdowns.",
        "5.  Step 4 - click the draft email cell, press Ctrl+C, paste into Outlook, and fine-tune before sending.",
    ]
    r = 6
    for s in steps:
        ws[f"C{r}"] = s
        ws[f"C{r}"].font = Font(name=FONT, size=11)
        ws[f"C{r}"].alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 22
        r += 1

    r += 1
    step_header(ws, f"B{r}:C{r}", "Tips")
    r += 1
    tips = [
        "*  You do not need to answer every question - the draft only includes paragraphs for the questions you have answered.",
        "*  To clear an answer, select the cell and press Delete.",
        "*  The wording lives on the 'Response library' and 'Common responses' tabs - edit the yellow columns there and the drafts update automatically.",
        "*  The draft is a starting point: add case-specific detail and always confirm the advice against the current Ikon guidance before sending.",
    ]
    for t in tips:
        ws[f"C{r}"] = t
        ws[f"C{r}"].font = Font(name=FONT, size=11)
        ws[f"C{r}"].alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 22
        r += 1

    r += 1
    step_header(ws, f"B{r}:C{r}", "Contact")
    r += 1
    ws[f"C{r}"] = "Principal wellbeing expenditure support  |  9264 4201  |  whsw.staffwellbeing@education.wa.edu.au"
    ws[f"C{r}"].font = Font(name=FONT, size=11, bold=True)
    ws.row_dimensions[r].height = 22
    r += 2
    note(ws, f"C{r}", "This information is correct as at July 2026. Always refer to Ikon for the most up-to-date information.")

    # ================================================== Email builder ==
    eb = wb.create_sheet("Email builder")
    eb.sheet_properties.tabColor = TEAL
    no_gridlines(eb)
    for col, w in {"A": 2, "B": 7, "C": 60, "D": 28, "E": 2, "F": 46, "G": 2}.items():
        eb.column_dimensions[col].width = w

    eb.row_dimensions[2].height = 34
    banner(eb, "B2:F2", "Principal wellbeing funding - email response builder", MAGENTA, size=18)
    eb.row_dimensions[3].height = 20
    banner(eb, "B3:F3", "Answer the questions below, then copy the draft email from Step 4", TEAL, size=11, bold=False)

    # ---- Step 1
    step_header(eb, "B5:F5", "STEP 1  ·  Query details")
    label(eb, "C6", "Greeting")
    input_cell(eb, "D6", "Good morning")
    note(eb, "F6", "Opens the email.")
    label(eb, "C7", "Recipient name (optional)")
    input_cell(eb, "D7")
    note(eb, "F7", 'Personalises the greeting, e.g. "Good morning Damien,"')
    label(eb, "C8", "Topic of the query (optional)")
    input_cell(eb, "D8")
    note(eb, "F8", 'Completes "Thank you for reaching out regarding ..." - e.g. "using the funding for physiotherapy"')

    # ---- Step 2
    step_header(eb, "B10:F10", "STEP 2  ·  Recurring query?  Pick a common response (optional)")
    label(eb, "C11", "Common response")
    eb.merge_cells("D11:F11")
    input_cell(eb, "D11", BUILD_OPTION)
    for coord in ("E11", "F11"):
        eb[coord].fill = fill(INPUT_YELLOW)
        eb[coord].border = box
    note(eb, "C12", f"Leave as '{BUILD_OPTION}' to build the reply from Step 3 instead.")

    # ---- Step 3
    step_header(eb, "B14:F14", "STEP 3  ·  Answer the framework questions (skip if you picked a common response)")
    note(eb, "C15", "Select Yes, No or Maybe. Unanswered questions are simply left out of the draft.")
    q_rows = {}
    r = 16
    for code, text in QUESTIONS + ROUTING:
        eb[f"B{r}"] = code
        eb[f"B{r}"].font = Font(name=FONT, size=10, bold=True, color=MAGENTA)
        eb[f"B{r}"].alignment = Alignment(vertical="center")
        label(eb, f"C{r}", text, bold=False)
        input_cell(eb, f"D{r}")
        eb[f"D{r}"].alignment = Alignment(horizontal="center", vertical="center")
        eb.row_dimensions[r].height = 30
        q_rows[code] = r
        if code == "Q6":
            r += 1
            note(eb, f"C{r}", "Extra routing questions - these add procurement / travel wording to the draft:")
            r += 1
        else:
            r += 1
    last_q_row = q_rows["Travel"]

    # ---- Result
    res_row = last_q_row + 2
    step_header(eb, f"B{res_row}:F{res_row}", "TRIAGE RESULT")
    band = res_row + 1
    eb.merge_cells(f"B{band}:F{band}")
    eb.row_dimensions[band].height = 30
    rc = eb[f"B{band}"]
    rc.font = Font(name=FONT, size=12, bold=True)
    rc.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    rc.border = box

    # ---- Step 4
    s4 = band + 2
    step_header(eb, f"B{s4}:F{s4}", "STEP 4  ·  Your draft email")
    note(eb, f"C{s4+1}", "Click the draft below, press Ctrl+C to copy, then paste into Outlook. Add case-specific detail and confirm against current Ikon guidance before sending.")
    d_first, d_last = s4 + 2, s4 + 26
    eb.merge_cells(f"B{d_first}:F{d_last}")
    dc = eb[f"B{d_first}"]
    dc.font = Font(name=FONT, size=10)
    dc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    dc.fill = fill(DRAFT_BG)
    for row in range(d_first, d_last + 1):
        for col in range(2, 7):
            eb.cell(row=row, column=col).border = box
            if eb.cell(row=row, column=col).fill.fgColor.rgb in (None, "00000000"):
                eb.cell(row=row, column=col).fill = fill(DRAFT_BG)

    # ================================================== Response library ==
    rl = wb.create_sheet("Response library")
    rl.sheet_properties.tabColor = MAGENTA
    no_gridlines(rl)
    for col, w in {"A": 10, "B": 60, "C": 10, "D": 100, "E": 14}.items():
        rl.column_dimensions[col].width = w
    rl.row_dimensions[1].height = 28
    banner(rl, "A1:E1", "Response library - modular paragraphs for the question-built draft", MAGENTA, size=13)
    note(rl, "A2", "Edit the yellow 'Draft email paragraph' column to change the wording. Do not edit column E (lookup key).")
    rl.merge_cells("A2:E2")
    headers = ["Code", "Question", "Answer", "Draft email paragraph - edit this yellow column", "Lookup key"]
    for i, h in enumerate(headers, 1):
        c = rl.cell(row=4, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    qtext = dict(QUESTIONS + ROUTING)
    for i, (code, ans, para) in enumerate(LIBRARY, start=5):
        rl.cell(row=i, column=1, value=code).font = Font(name=FONT, size=10, bold=True, color=MAGENTA)
        qc = rl.cell(row=i, column=2, value=qtext[code])
        qc.font = Font(name=FONT, size=9, color=GREY_TEXT)
        qc.alignment = Alignment(vertical="top", wrap_text=True)
        rl.cell(row=i, column=3, value=ans).font = Font(name=FONT, size=10)
        pc = rl.cell(row=i, column=4, value=para)
        pc.fill = fill(INPUT_YELLOW)
        pc.border = box
        pc.font = Font(name=FONT, size=10)
        pc.alignment = Alignment(vertical="top", wrap_text=True)
        kc = rl.cell(row=i, column=5, value=f'=A{i}&"|"&C{i}')
        kc.font = Font(name=FONT, size=9, color=GREY_TEXT)
        rl.row_dimensions[i].height = max(30, 14 * (len(para) // 95 + 1) + 8)
    lib_last = 4 + len(LIBRARY)

    # ================================================== Common responses ==
    cr = wb.create_sheet("Common responses")
    cr.sheet_properties.tabColor = MAGENTA
    no_gridlines(cr)
    for col, w in {"A": 55, "B": 120}.items():
        cr.column_dimensions[col].width = w
    cr.row_dimensions[1].height = 28
    banner(cr, "A1:B1", "Common responses - full reply bodies for recurring queries", MAGENTA, size=13)
    note(cr, "A2", "These appear in the Step 2 dropdown on the Email builder tab. Edit the yellow text (or the option names) and the tool updates automatically. The greeting, thank-you line and sign-off are added by the builder - only write the body here.")
    cr.merge_cells("A2:B2")
    cr.row_dimensions[2].height = 30
    for i, h in enumerate(["Response option (appears in the dropdown)", "Draft email body - edit this yellow column"], 1):
        c = cr.cell(row=4, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, (name, body) in enumerate(PRESETS, start=5):
        nc = cr.cell(row=i, column=1, value=name)
        nc.font = Font(name=FONT, size=10, bold=True)
        nc.alignment = Alignment(vertical="top", wrap_text=True)
        bc = cr.cell(row=i, column=2, value=body)
        bc.fill = fill(INPUT_YELLOW)
        bc.border = box
        bc.font = Font(name=FONT, size=10)
        bc.alignment = Alignment(vertical="top", wrap_text=True)
        lines = sum(len(p) // 115 + 1 for p in body.split("\n\n")) + body.count("\n\n")
        cr.row_dimensions[i].height = max(60, 14 * lines + 10)

    # ================================================== Reference links ==
    lk = wb.create_sheet("Reference links")
    lk.sheet_properties.tabColor = TEAL
    no_gridlines(lk)
    for col, w in {"A": 2, "B": 62, "C": 95}.items():
        lk.column_dimensions[col].width = w
    lk.row_dimensions[2].height = 28
    banner(lk, "B2:C2", "Reference links", MAGENTA, size=13)
    for i, h in enumerate(["Topic", "Link"], 2):
        c = lk.cell(row=4, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = fill(TEAL)
    for i, (topic, url) in enumerate(LINKS, start=5):
        lk.cell(row=i, column=2, value=topic).font = Font(name=FONT, size=10)
        c = lk.cell(row=i, column=3, value=url)
        c.hyperlink = url
        c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        lk.row_dimensions[i].height = 18

    # ================================================== Lists (hidden) ==
    ls = wb.create_sheet("Lists")
    ls["A1"] = "Greetings"
    for i, g in enumerate(["Good morning", "Good afternoon", "Good evening", "Hi"], 2):
        ls[f"A{i}"] = g
    ls["B1"] = "Answers"
    for i, a in enumerate(["Yes", "No", "Maybe"], 2):
        ls[f"B{i}"] = a
    ls["C1"] = "Response types"
    ls["C2"] = BUILD_OPTION
    for i in range(len(PRESETS)):
        ls[f"C{3+i}"] = f"=IF('Common responses'!A{5+i}=\"\",\"\",'Common responses'!A{5+i})"

    # ---- logic
    ebq = lambda code: f"'Email builder'!$D${q_rows[code]}"
    ls["F1"] = "Logic"
    ls["F2"] = ('=TRIM(IF(\'Email builder\'!$D$6="","Hi",\'Email builder\'!$D$6)'
                '&" "&TRIM(\'Email builder\'!$D$7))&","')
    ls["F3"] = ('=IF(TRIM(\'Email builder\'!$D$8)="",'
                '"Thank you for reaching out regarding the principal wellbeing targeted funding.",'
                '"Thank you for reaching out regarding "&TRIM(\'Email builder\'!$D$8)&".")')
    codes = ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Panel", "Travel"]
    for i, code in enumerate(codes):
        row = 2 + i
        ls[f"G{row}"] = f"={ebq(code)}"
        ls[f"H{row}"] = (f'=IF(G{row}="","",IFERROR(T(INDEX(\'Response library\'!$D$5:$D${lib_last},'
                         f'MATCH("{code}|"&G{row},\'Response library\'!$E$5:$E${lib_last},0))),""))')
    ls["F5"] = "=COUNTIF(G2:G7,\"?*\")"
    ls["F6"] = ('=IF(F5=0,"",IF(COUNTIF(G2:G7,"No")>0,'
                '"To answer briefly: based on the decision-making framework, this expenditure does not appear to be appropriate. Please see the detail below.",'
                'IF(COUNTIF(G2:G7,"Yes")=6,'
                '"To answer briefly: based on the decision-making framework, this expenditure is likely to be appropriate. Please see the detail below.",'
                '"To answer briefly: we need a little more information before this expenditure can be confirmed as appropriate. Please see the detail below.")))')
    ls["F7"] = (f'=IF(F5=0,"",_xlfn.TEXTJOIN({NL2},TRUE,F6,"To answer more thoroughly:",H2:H9))')
    ls["F8"] = "='Email builder'!$D$11"
    ls["F9"] = (f'=IF(OR(F8="",F8="{BUILD_OPTION}"),"",'
                f"IFERROR(T(INDEX('Common responses'!$B$5:$B${4+len(PRESETS)},"
                f"MATCH(F8,'Common responses'!$A$5:$A${4+len(PRESETS)},0))),\"\"))")
    ls["F10"] = f'=IF(F9<>"",F9,IF(F5=0,"{PLACEHOLDER}",F7))'
    ls["F11"] = f'="{CONTACT}"'
    ls["F12"] = f'=F2&{NL2}&F3&{NL2}&F10&{NL2}&F11&{NL2}&"Kind regards,"'
    ls["F13"] = ('=IF(F9<>"","Common response selected: "&F8,'
                 'IF(F5=0,"Waiting for answers - pick a common response (Step 2) or answer the questions (Step 3).",'
                 'IF(F5<6,"Keep going - "&F5&" of 6 framework questions answered. The draft below updates as you go.",'
                 'IF(COUNTIF(G2:G7,"No")>0,"Likely NOT appropriate - recommend reconsidering or contacting the team.",'
                 'IF(COUNTIF(G2:G7,"Yes")=6,"Likely appropriate to proceed.",'
                 '"Needs more information - recommend contacting Principal wellbeing expenditure support.")))))')
    ls.sheet_state = "hidden"

    # wire builder result + draft
    rc.value = "=Lists!F13"
    dc.value = "=Lists!F12"

    # ---- data validation
    dv_greet = DataValidation(type="list", formula1="'Lists'!$A$2:$A$5", allow_blank=True)
    dv_ans = DataValidation(type="list", formula1="'Lists'!$B$2:$B$4", allow_blank=True,
                            errorTitle="Pick from the list", error="Please choose Yes, No or Maybe.")
    dv_resp = DataValidation(type="list", formula1=f"'Lists'!$C$2:$C${2+len(PRESETS)}", allow_blank=True)
    eb.add_data_validation(dv_greet)
    eb.add_data_validation(dv_ans)
    eb.add_data_validation(dv_resp)
    dv_greet.add("D6")
    dv_resp.add("D11")
    for code in codes:
        dv_ans.add(f"D{q_rows[code]}")

    # ---- conditional formatting
    ans_range = f"D{q_rows['Q1']}:D{q_rows['Travel']}"
    eb.conditional_formatting.add(ans_range, CellIsRule(
        operator="equal", formula=['"Yes"'], fill=fill(GREEN_FILL),
        font=Font(name=FONT, size=10, bold=True, color=GREEN_FONT)))
    eb.conditional_formatting.add(ans_range, CellIsRule(
        operator="equal", formula=['"No"'], fill=fill(RED_FILL),
        font=Font(name=FONT, size=10, bold=True, color=RED_FONT)))
    eb.conditional_formatting.add(ans_range, CellIsRule(
        operator="equal", formula=['"Maybe"'], fill=fill(AMBER_FILL),
        font=Font(name=FONT, size=10, bold=True, color=AMBER_FONT)))
    band_cell = f"B{band}"
    eb.conditional_formatting.add(band_cell, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("Likely appropriate",{band_cell}))'],
        fill=fill(GREEN_FILL), font=Font(name=FONT, size=12, bold=True, color=GREEN_FONT)))
    eb.conditional_formatting.add(band_cell, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("NOT appropriate",{band_cell}))'],
        fill=fill(RED_FILL), font=Font(name=FONT, size=12, bold=True, color=RED_FONT)))
    eb.conditional_formatting.add(band_cell, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("Common response",{band_cell}))'],
        fill=fill(TEAL_TINT), font=Font(name=FONT, size=12, bold=True, color=TEAL)))
    eb.conditional_formatting.add(band_cell, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("more information",{band_cell}))'],
        fill=fill(AMBER_FILL), font=Font(name=FONT, size=12, bold=True, color=AMBER_FONT)))
    eb.conditional_formatting.add(band_cell, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("Keep going",{band_cell}))'],
        fill=fill(AMBER_FILL), font=Font(name=FONT, size=12, bold=True, color=AMBER_FONT)))

    wb.save(path)
    print(f"Saved {path}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "Principal wellbeing email response builder.xlsx")
