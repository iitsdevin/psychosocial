#!/usr/bin/env python3
"""
Build a non-macro version of the Level 1 PHIR Response document with
interactive checkbox content controls for PEEPO factors embedded
directly in the 'Outline any relevant work or non-work factors' cell.
"""

import os
import copy
from docx import Document
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from lxml import etree

UPLOAD_DIR = "/root/.claude/uploads/9ca4ee02-85b9-52bf-99ea-c1789f8f94da"
SOURCE_DOCX = os.path.join(UPLOAD_DIR, "13f36241-Level1PHIRResponse_9.DOCX")
OUTPUT_DOCX = "/home/user/psychosocial/Level1_PHIR_Response_WithCheckboxes.docx"

PEEPO_CATEGORIES = {
    "People (Human Factors)": [
        "Inadequate communication between staff",
        "Unreasonable behaviour by worker(s)",
        "Unreasonable behaviour by manager(s) or supervisor(s)",
        "Interpersonal conflict (unresolved or escalating)",
        "Fatigue or burnout",
        "Mental stress or psychological distress",
        "Exposure to traumatic event or material",
        "Personal issues affecting work capacity",
        "Incivility or inappropriate behaviour tolerated by peers",
    ],
    "Equipment (Resources & Equipment)": [
        "Insufficient staffing levels or resources to meet work demands",
        "Inadequate or unavailable tools and resources to do the job",
        "Insufficient training on psychosocial hazard identification or management",
        "Insufficient training for job demands or new tasks",
        "Lack of constructive feedback or performance support",
        "Inadequate access to employee assistance or support services",
        "Lack of information about reporting processes or support pathways",
        "Insufficient or inappropriate supervision",
        "Inadequate de-escalation tools or conflict resolution resources",
    ],
    "Environment (Physical & Workplace)": [
        "Adverse environmental conditions (noise, temperature, air quality)",
        "Poorly maintained or unclean facilities",
        "Workplace layout limiting supervision or support access",
        "Remote or isolated work location",
        "Natural events restricting travel or creating uncertainty",
        "Limited access to communication technology",
        "Blurring of boundaries between work and home life",
        "Workplace culture tolerating unreasonable behaviour",
        "Lack of diversity, respect, or inclusion in the workplace",
        "Limited opportunities for social interaction during work",
    ],
    "Processes (Tasks & Procedures)": [
        "No or inadequate policies for managing psychosocial hazards",
        "Policies or procedures not adhered to",
        "Procedures that cannot be applied as written or lack flexibility",
        "No or inadequate reporting and complaints processes",
        "Investigation or appeals process not affording natural justice",
        "No or inadequate incident response or investigation procedures",
        "Procedures that discriminate or lack fairness",
        "Inadequate rostering, hours of work, or fatigue management procedures",
        "Inadequate change management procedures",
        "No mechanism for impartially addressing behaviour by senior management",
    ],
    "Organisation (System & Policy)": [
        "Poor leadership practices or management style",
        "Limited management accountability for psychosocial hazards",
        "Unclear or conflicting roles, responsibilities, or expectations",
        "Excessive work demands (cognitive, emotional, or physical)",
        "Workplace culture",
        "Multiple interacting psychosocial hazards not considered together",
        "Low job control or autonomy",
        "Lack of recognition or reward",
        "Organisational injustice (unfairness, bias, inconsistency)",
        "Inadequate consultation with workers on matters affecting them",
        "Poorly managed organisational change",
        "Insecure work arrangements",
        "Inadequate monitoring and review of existing controls",
    ],
}


def make_checkbox_sdt(label_text, font_size=8):
    """Create a structured document tag (SDT) checkbox with a label."""

    # We need the w14 namespace for checkbox controls
    W14_NS = 'http://schemas.microsoft.com/office/word/2010/wordml'
    W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

    # Build the paragraph containing: [checkbox SDT] [space] [label text]
    para = OxmlElement('w:p')

    # Paragraph properties - small spacing
    pPr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:after'), '20')
    spacing.set(qn('w:before'), '0')
    spacing.set(qn('w:line'), '240')
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)
    para.append(pPr)

    # SDT for the checkbox
    sdt = OxmlElement('w:sdt')

    # SDT Properties
    sdtPr = OxmlElement('w:sdtPr')

    # Checkbox element (w14 namespace)
    checkbox = etree.SubElement(sdtPr, f'{{{W14_NS}}}checkbox')
    checked = etree.SubElement(checkbox, f'{{{W14_NS}}}checked')
    checked.set(f'{{{W14_NS}}}val', '0')
    checkedState = etree.SubElement(checkbox, f'{{{W14_NS}}}checkedState')
    checkedState.set(f'{{{W14_NS}}}val', '2612')
    checkedState.set(f'{{{W14_NS}}}font', 'MS Gothic')
    uncheckedState = etree.SubElement(checkbox, f'{{{W14_NS}}}uncheckedState')
    uncheckedState.set(f'{{{W14_NS}}}val', '2610')
    uncheckedState.set(f'{{{W14_NS}}}font', 'MS Gothic')

    sdt.append(sdtPr)

    # SDT Content
    sdtContent = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'MS Gothic')
    rFonts.set(qn('w:eastAsia'), 'MS Gothic')
    rFonts.set(qn('w:hAnsi'), 'MS Gothic')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(font_size * 2))
    rPr.append(sz)
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(font_size * 2))
    rPr.append(szCs)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = '☐'  # ☐ ballot box
    run.append(t)
    sdtContent.append(run)
    sdt.append(sdtContent)

    para.append(sdt)

    # Space + label text
    label_run = OxmlElement('w:r')
    label_rPr = OxmlElement('w:rPr')
    label_sz = OxmlElement('w:sz')
    label_sz.set(qn('w:val'), str(font_size * 2))
    label_rPr.append(label_sz)
    label_szCs = OxmlElement('w:szCs')
    label_szCs.set(qn('w:val'), str(font_size * 2))
    label_rPr.append(label_szCs)
    label_run.append(label_rPr)
    label_t = OxmlElement('w:t')
    label_t.set(qn('xml:space'), 'preserve')
    label_t.text = ' ' + label_text
    label_run.append(label_t)
    para.append(label_run)

    return para


def make_category_heading(text, font_size=9):
    """Create a bold category heading paragraph."""
    para = OxmlElement('w:p')

    pPr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), '120')
    spacing.set(qn('w:after'), '40')
    spacing.set(qn('w:line'), '240')
    spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)
    para.append(pPr)

    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    bold = OxmlElement('w:b')
    rPr.append(bold)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(font_size * 2))
    rPr.append(sz)
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(font_size * 2))
    rPr.append(szCs)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '003366')
    rPr.append(color)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    para.append(run)

    return para


def find_target_cell(doc):
    """Find the cell after 'Outline any relevant work or non-work factors'."""
    for table in doc.tables:
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if 'Outline any relevant work or non-work factors' in cell.text:
                    # Target is the next row, first column
                    try:
                        target = table.cell(r_idx + 1, 0)
                        return target
                    except IndexError:
                        try:
                            target = table.cell(r_idx + 1, c_idx)
                            return target
                        except IndexError:
                            pass
    return None


def main():
    print("Opening source document...")
    doc = Document(SOURCE_DOCX)

    # Register the w14 namespace
    W14_NS = 'http://schemas.microsoft.com/office/word/2010/wordml'
    etree.register_namespace('w14', W14_NS)

    print("Finding target cell...")
    target_cell = find_target_cell(doc)

    if target_cell is None:
        print("ERROR: Could not find target cell!")
        return

    print(f"Found target cell with text: '{target_cell.text[:50]}...'")

    # Clear existing content in the target cell
    tc = target_cell._tc
    # Remove all existing paragraphs except keep the cell structure
    for p in list(tc.findall(qn('w:p'))):
        tc.remove(p)

    # Add instruction text
    instruction = OxmlElement('w:p')
    iPr = OxmlElement('w:pPr')
    iSpacing = OxmlElement('w:spacing')
    iSpacing.set(qn('w:after'), '60')
    iPr.append(iSpacing)
    instruction.append(iPr)

    iRun = OxmlElement('w:r')
    iRPr = OxmlElement('w:rPr')
    iItalic = OxmlElement('w:i')
    iRPr.append(iItalic)
    iSz = OxmlElement('w:sz')
    iSz.set(qn('w:val'), '16')
    iRPr.append(iSz)
    iSzCs = OxmlElement('w:szCs')
    iSzCs.set(qn('w:val'), '16')
    iRPr.append(iSzCs)
    iColor = OxmlElement('w:color')
    iColor.set(qn('w:val'), '666666')
    iRPr.append(iColor)
    iRun.append(iRPr)
    iT = OxmlElement('w:t')
    iT.text = 'Click the boxes below to select applicable work factors:'
    iRun.append(iT)
    instruction.append(iRun)
    tc.append(instruction)

    # Add each category with its checkboxes
    for cat_name, factors in PEEPO_CATEGORIES.items():
        # Category heading
        heading = make_category_heading(cat_name)
        tc.append(heading)

        # Factor checkboxes
        for factor in factors:
            cb_para = make_checkbox_sdt(factor)
            tc.append(cb_para)

    # Ensure w14 namespace is declared in the document root (only if not already present)
    root = doc.element
    existing_ns = root.nsmap
    if 'w14' not in existing_ns:
        root.attrib['{http://www.w3.org/2000/xmlns/}w14'] = W14_NS

    print(f"Saving to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print("Done!")
    print()
    print("To use: Open in Word, click any checkbox to tick/untick it.")
    print("No macros needed - checkboxes work natively in Word.")


if __name__ == "__main__":
    main()
